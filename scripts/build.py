"""Generate the static career package from shared facts and checked sample inputs."""
from pathlib import Path
from html import escape as esc
import argparse, json, shutil, math, re

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'dist'
PROFILE = json.loads((ROOT / 'content/profile.json').read_text(encoding='utf-8'))
BASE = ''

def table(headers, rows, caption=''):
    head = ''.join(f'<th scope="col">{esc(str(x))}</th>' for x in headers)
    body = ''.join('<tr>'+''.join(f'<td>{esc(str(v))}</td>' for v in row)+'</tr>' for row in rows)
    return f'<div class="table-wrap" role="region" aria-label="{esc(caption or "Data table")}" tabindex="0"><table>{"<caption>"+esc(caption)+"</caption>" if caption else ""}<thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>'

def ul(items): return '<ul>'+''.join('<li>'+esc(x)+'</li>' for x in items)+'</ul>'
def p(text): return '<p>'+esc(text)+'</p>'
def btn(text, href, secondary=False): return f'<a class="button{" secondary" if secondary else ""}" href="{esc(href)}">{esc(text)}</a>'

def svg_shell(body, label, dark=False):
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 600 360" role="img" aria-label="{esc(label)}"><title>{esc(label)}</title><rect width="600" height="360" fill="{"#183b43" if dark else "#e7efe7"}"/>{body}</svg>'

POINTS = {
  'lighting': [(x,y) for y in [105,180,255] for x in [145,250,355,460]],
  'outlets': [(x,y) for x,y in [(95,100),(95,170),(95,250),(185,300),(280,300),(380,300),(485,300),(515,100),(515,175),(515,250),(185,65),(315,65),(445,65)]]
}

def plan_svg(dark=False):
    color='#c6e4d6' if dark else '#2d6453'
    b=f'<g fill="none" stroke="{color}" stroke-width="2"><path d="M85 55H525V310H85Z M305 55V135 M305 225V310"/><path d="M305 135A90 90 0 0 1 395 225" stroke-dasharray="4 5"/><path d="M305 225H395"/>'
    for x,y in POINTS['lighting']:b+=f'<circle cx="{x}" cy="{y}" r="12"/><path d="M{x-8} {y-8}l16 16m-16 0l16 -16"/>'
    for x,y in POINTS['outlets']:b+=f'<rect x="{x-6}" y="{y-6}" width="12" height="12" fill="{color}"/>'
    b+=f'</g><g font-family="Arial,sans-serif" fill="{color}" font-size="13"><text x="85" y="30">E-SAMPLE-01 / DEVICE COUNT</text><text x="85" y="340">12 LIGHT POINTS</text><text x="275" y="340">13 OUTLET LOCATIONS</text><text x="85" y="290">ZONE A</text><text x="355" y="290">ZONE B</text></g>'
    return svg_shell(b,'Illustrative sample layout with twelve light points and thirteen outlet locations. Not to scale.',dark)

def load_svg():
    b='<g font-family="Arial,sans-serif" fill="#173c43"><text x="45" y="42" font-size="15">CONNECTED LOAD / DEMONSTRATION</text>'
    for label,value,y,width in [('Lighting',0.384,90,76.8),('Workstations',1.32,155,264),('Network equipment',0.09,220,18)]:
        b+=f'<text x="45" y="{y}" font-size="16">{label}</text><rect x="45" y="{y+12}" width="{width}" height="17" rx="3" fill="#32725a"/><text x="465" y="{y+27}" font-size="18">{value:.3f} kVA</text>'
    b+='<text x="45" y="322" font-size="22">Total 1.794 kVA / 7.80 A at 230 V</text></g>'
    return svg_shell(b,'Connected apparent load: lighting 0.384 kVA, workstations 1.320 kVA, network 0.090 kVA; total 1.794 kVA.')

def equipment_svg():
    b='<g font-family="Arial,sans-serif" fill="#173c43"><text x="42" y="40" font-size="15">24 V DC / STEADY-STATE CAPACITY</text>'
    for x,label,amps,state in [(42,'A',1.5,'Below target'),(224,'B',2.5,'Meets target'),(406,'C',1.7,'Derated below')]:
        b+=f'<rect x="{x}" y="75" width="152" height="205" rx="8" fill="white" stroke="#a8bcb0"/><text x="{x+18}" y="110" font-size="17">OPTION {label}</text><text x="{x+18}" y="158" font-size="30">{amps:.1f} A</text><text x="{x+18}" y="197" font-size="12">{state}</text><path d="M{x+18} 225h115m-115 10h115m-115 10h115" stroke="#a8bcb0"/>'
    b+='<text x="42" y="325" font-size="19">Comparison target: 1.875 A / 45 W</text></g>'
    return svg_shell(b,'Sample power supply options A 1.5 A, B 2.5 A, C 1.7 A after assumed derating. Target is 1.875 A.')

def records_svg():
    b='<g font-family="Arial,sans-serif" fill="#173c43"><text x="42" y="38" font-size="15">DOCUMENT CONTROL / REVIEW STATUS</text>'
    for i,(title,status,rev) in enumerate([('Equipment register','Current','C'),('Cable schedule','Required C / received B','B'),('Test record','Missing attachment','A')]):
        y=70+i*80;b+=f'<rect x="42" y="{y}" width="516" height="63" rx="4" fill="white" stroke="#a8bcb0"/><text x="60" y="{y+25}" font-size="17">{title}</text><text x="60" y="{y+46}" font-size="12">{status}</text><text x="468" y="{y+37}" font-size="15">RCVD {rev}</text>'
    b+='</g>';return svg_shell(b,'Sample document register: equipment register received C, cable schedule received B with C required, test record missing attachment.')

CASES = [
 {"slug":"takeoff","title":"Sample Electrical Estimating Take-Off","classification":"Sample project","intro":"Device counts, route quantities and scope clarifications organized from an illustrative office layout.","context":"The sample layout contains two zones, 12 lighting points and 13 outlet locations. Three horizontal route lengths and a separate allowance are assumed inputs, not measurements from construction drawings.","objective":"Produce a symbol count, route-length build-up and clarification register that another estimator can check.","approach":"Assign stable point IDs, count each symbol once and reconcile the total with the point register. Keep assumed route lengths separate from cable quantities and itemize rises and termination allowances. Preserve the quantity basis without adding a blanket percentage.","tools":"Point register, arithmetic checks, SVG layout and quantity schedule. GroundPlan is used in my separate ongoing training; these route lengths were supplied as sample inputs.","deliverables":"Illustrative symbol layout, 25-row point register, quantity schedule, route build-up and scope questions.","checks":"12 unique lighting points and 13 unique outlet locations match the layout. Assumed route segments total 39 m; the separate 6 m rise and termination allowance gives a 45 m working route quantity. No blanket percentage is added.","result":"Each quantity has a stated basis, with unresolved circuit, route and specification questions recorded for clarification.","caption":"Illustrative device layout: 12 lighting points and 13 outlet locations. Not to scale.","limit":"Route segments and the 6 m allowance are sample assumptions. SVG coordinates locate symbols only and do not represent metres. The 45 m result is a route quantity, not a cable or conductor quantity. Cable sizes, circuit allocation, installation requirements, labour and prices remain outside this sample.","category":"Electrical estimating","visual":plan_svg},
 {"slug":"schedule-review","title":"Sample Electrical Load Calculation","classification":"Sample calculation","intro":"A connected-load schedule using assumed equipment ratings, consistent VA units and an explicit supply-voltage assumption.","context":"The sample uses assumed apparent-power ratings for lighting, workstations and network equipment. Each quantity and VA rating is shown in the schedule.","objective":"Create a traceable connected-load schedule and identify the information still needed for circuit design.","approach":"Multiply each quantity by its assumed apparent power in VA and sum the entries. Divide total VA by an assumed single-phase 230 V supply to illustrate an aggregate current budget. No diversity, energy-use estimate or protective-device selection is applied.","tools":"Load schedule, arithmetic checks and a diagram of the connected apparent-power totals.","deliverables":"Connected-load schedule, current calculation, assumptions and design review questions.","checks":"12 lights at 32 VA = 384 VA; six workstations at 220 VA = 1,320 VA; one network load at 90 VA = 90 VA. Total = 1,794 VA (1.794 kVA). At 230 V, the illustrative current budget is 7.80 A.","result":"The schedule links each subtotal to its device count and assumed VA rating. VA and current describe connected load; energy consumption would require active power in W and an operating period.","caption":"Sample connected apparent load: 0.384 + 1.320 + 0.090 = 1.794 kVA.","limit":"All equipment ratings and the 230 V supply are assumed. The simple VA sum is a preliminary current budget, without power-factor or phase-angle assessment. Actual ratings, simultaneity, harmonics, starting current, voltage drop, fault level and installation conditions require review before cable or protection selection. This is not an issued design.","category":"Electrical documentation","visual":load_svg},
 {"slug":"equipment-evaluation","title":"Sample Equipment Evaluation","classification":"Sample technical analysis","intro":"A 24 V DC power-supply comparison using a stated current budget, example reserve and assumed available output ratings.","context":"Four device groups have assumed steady-state currents at 24 V DC. Options A, B and C are illustrative specifications rather than identified products or supplier quotations.","objective":"Compare output capacity and prepare supplier questions before a quotation recommendation.","approach":"Sum the steady-state currents and add a stated 25% example reserve. Compare the target with each option’s assumed available current. Option C includes an assumed temperature derating; a numerical capacity check does not establish suitability for purchase.","tools":"Requirement schedule, comparison matrix, browser calculator and arithmetic checks.","deliverables":"Device current budget, three-option comparison and supplier clarification checklist.","checks":"The four currents total 1.500 A. A 25% reserve gives 1.875 A, equivalent to 45 W at 24 V DC. Option B’s assumed 2.5 A available capacity meets that target. Option C is compared at its assumed derated 1.7 A, rather than its nominal 2.0 A.","result":"Option B passes the sample steady-state capacity check. A product recommendation would also require manufacturer data and the actual application conditions.","caption":"Illustrative available capacities at 24 V DC. Comparison target: 1.875 A / 45 W.","limit":"Device currents, supply ratings and the derating are assumed. The 25% reserve is an example allowance, not a code requirement. Inrush, overload response, input supply, ambient conditions, dimensions, approvals, availability and warranty need manufacturer confirmation. No purchase or sales outcome is represented.","category":"Technical sales support","visual":equipment_svg},
 {"slug":"document-control","title":"Sample Engineering Document Register","classification":"Sample documentation","intro":"Revision tracking, attachment checks and review status organized in a five-record engineering register.","context":"The register uses sample document IDs, revisions, statuses and role assignments to show checks commonly needed when organizing an engineering handover.","objective":"Prepare a document register and exception log with a clear reason for each outstanding item.","approach":"Assign a stable record ID and compare required and received revisions. Check attachment presence and review status separately. Keep revision mismatches, missing attachments and pending reviews on hold until the relevant requirements are satisfied.","tools":"Document register, structured records, exception table and status checks.","deliverables":"Five-record register, three-item exception log, naming example and handover checks.","checks":"Each document has one unique record. Two records satisfy the sample transmittal checks. Three remain on hold: a cable-schedule revision mismatch, a missing test attachment and a pending manual review.","result":"The register shows the outstanding requirement and next action for each record. A ready status refers only to the listed document checks.","caption":"First three sample records. The cable schedule was received at revision B; revision C is required.","limit":"Document IDs, statuses and role assignments are illustrative. A ready-for-transmittal status is not technical approval of the underlying document or the complete package. The register reproduces no employer records or facility identifiers.","category":"Engineering operations","visual":records_svg}
]

SAMPLE_PROVENANCE = 'Prepared with AI assistance using illustrative inputs to demonstrate engineering methods. Samples are separate from professional client deliveries, issued designs and procurement approvals.'
SAMPLE_BASIS = {
 "takeoff": "Device symbols and SVG coordinates are constructed for the sample. Route lengths of 18, 12 and 9 m and a 6 m rise/termination allowance are assumed, not measured from this unscaled layout.",
 "schedule-review": "Quantities, equipment apparent-power ratings in VA and a single-phase 230 V supply are assumed. No operating hours, active-power ratings or measured energy use are provided.",
 "equipment-evaluation": "All device currents and option capacities are assumed at 24 V DC. The 25% reserve and Option C derating are example assumptions, not manufacturer or code requirements.",
 "document-control": "Document IDs, required/received revisions, attachments, review statuses and owners are sample records. Status indicates register checks only, not engineering approval."
}

DEMO = {
 'classification':'Engineering work samples with illustrative inputs.',
 'points':[{'id':f'{prefix}{i:02d}','type':kind,'x':x,'y':y} for kind,prefix in [('lighting','L'),('outlets','P')] for i,(x,y) in enumerate(POINTS[kind],1)],
 'route':{'segments_m':[18,12,9],'rise_and_termination_allowance_m':6},
 'loads':[{'item':'LED luminaires','quantity':12,'unit_va':32},{'item':'Workstations','quantity':6,'unit_va':220},{'item':'Network equipment','quantity':1,'unit_va':90}],
 'supply':{'voltage_v':24,'load_a':[0.30,0.80,0.15,0.25],'reserve_fraction':0.25,'options':[{'name':'A','nominal_a':1.5,'available_a':1.5},{'name':'B','nominal_a':2.5,'available_a':2.5},{'name':'C','nominal_a':2.0,'available_a':1.7}]},
 'records':[
  {'id':'DOC-001','name':'Equipment register','required':'C','received':'C','review':'Reviewed','attachment':True,'owner':'Document coordinator','next':'File current register'},
  {'id':'DOC-002','name':'Cable schedule','required':'C','received':'B','review':'Pending','attachment':True,'owner':'Discipline engineer','next':'Provide revision C and review'},
  {'id':'DOC-003','name':'Test record','required':'A','received':'A','review':'Pending','attachment':False,'owner':'Site engineer','next':'Attach signed record for review'},
  {'id':'DOC-004','name':'Equipment manual','required':'B','received':'B','review':'Pending','attachment':True,'owner':'Discipline engineer','next':'Complete manual review'},
  {'id':'DOC-005','name':'Drawing index','required':'C','received':'C','review':'Reviewed','attachment':True,'owner':'Document coordinator','next':'Include in transmittal'}]
}

def record_ready(r): return r['required']==r['received'] and r['review']=='Reviewed' and r['attachment']

def details(slug):
    if slug=='takeoff':
        return table(['Item','Basis','Quantity','Unit'],[['Lighting points','Counted symbols',12,'each'],['Outlet locations','Counted symbols',13,'each'],['Horizontal route','Assumed: 18 + 12 + 9',39,'m'],['Rises / terminations','Explicit assumed allowance',6,'m'],['Working route quantity','39 + 6',45,'m']], 'Takeoff summary')+p('Open RFIs: confirm circuit allocation and cable specification; confirm actual route and vertical dimensions; confirm whether termination slack is already included in the project measurement rules.')
    if slug=='schedule-review':
        rows=[[x['item'],x['quantity'],x['unit_va'],x['quantity']*x['unit_va']] for x in DEMO['loads']]
        return table(['Equipment','Quantity','VA per item','Connected VA'],rows+[['Total','','',1794]],'Connected apparent load')+p('Calculation: 1,794 VA ÷ 230 V = 7.80 A. This is a connected apparent-power budget, not energy consumption in kWh. Active input power and operating hours would be needed for an energy calculation; demand assessment and circuit design remain separate.')
    if slug=='equipment-evaluation':
        rows=[['Controller',0.30],['Actuators',0.80],['Sensors',0.15],['Auxiliaries',0.25]]
        return table(['Device group','Assumed current at 24 V DC (A)'],rows,'Steady-state load inputs')+table(['Sample option','Nominal A','Available A','Against 1.875 A'],[['A',1.5,1.5,'Below target'],['B',2.5,2.5,'Meets numerical target'],['C',2.0,1.7,'Below target after assumed derating']],'Quotation comparison')+'''<div class="calculator"><h3>Explore the capacity requirement</h3><p class="small">Change the assumed steady-state load or example reserve. Output is a capacity comparison only.</p><form id="supply-form"><div class="fields"><label for="load">Total load at 24 V DC (A)<input id="load" name="load" type="number" min="0" max="20" step="0.01" value="1.5" required></label><label for="reserve">Additional reserve (%)<input id="reserve" name="reserve" type="number" min="0" max="100" step="1" value="25" required></label></div><p class="result" id="supply-result" aria-live="polite">Required capacity: 1.875 A at 24 V DC (45.0 W). Option B meets this numerical check. Supplier confirmation remains required.</p></form></div>'''+p('Supplier questions: confirm continuous output at the actual enclosure temperature, transient load response, input range, dimensions and mounting, protection characteristics, lead time and quotation exclusions. No price comparison is invented.')
    return table(['Record','Required / received','Review','Next action'],[[r['name'],r['required']+' / '+r['received'], 'Ready for transmittal' if record_ready(r) else 'Hold',r['next']] for r in DEMO['records']],'Handover register')+p('Naming example: DEMO_DOC-002_Cable-Schedule_Rev-C. Preserve the original received file, record its revision, and create a reviewed issue without silently overwriting the earlier record. Only the authorized reviewer can approve technical content.')

def page(filename,title,description,body,current='',robots=True):
    nav=''.join(f'<a href="{path}"'+(' aria-current="page"' if current==key else '')+f'>{label}</a>' for key,label,path in [('experience','Experience','experience.html'),('expertise','Skills','engineering.html'),('projects','Work samples','projects.html'),('resumes','Résumé','resumes.html'),('contact','Contact','contact.html')])
    canonical=f'<link rel="canonical" href="{BASE}/{filename if filename!="index.html" else ""}"><meta property="og:url" content="{BASE}/{filename if filename!="index.html" else ""}">' if BASE else '<meta name="robots" content="noindex, nofollow">'
    out=f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc(title)} | Eric John Cruz</title><meta name="description" content="{esc(description)}"><meta name="theme-color" content="#122c35"><meta property="og:type" content="website"><meta property="og:title" content="{esc(title)} | Eric John Cruz"><meta property="og:description" content="{esc(description)}">{canonical}<link rel="icon" href="assets/favicon.svg" type="image/svg+xml"><link rel="stylesheet" href="assets/site.css"><script src="assets/site.js" defer></script></head><body><a class="skip" href="#main">Skip to content</a><header class="header"><div class="wrap"><a class="brand" href="index.html"><span class="brand-mark" aria-hidden="true">EC</span><span>Eric John Cruz<small>Electrical Engineering &amp; Technical Support</small></span></a><button class="menu" aria-expanded="false" aria-controls="primary-nav" type="button">Menu</button><nav class="nav" id="primary-nav" aria-label="Primary">{nav}</nav></div></header><main class="wrap" id="main">{body}</main><footer class="footer"><div class="wrap"><p>Eric John Cruz<br>Puerto Princesa City, Philippines</p><p><a href="mailto:{PROFILE['email']}">{PROFILE['email']}</a><br><a href="{PROFILE['linkedin']}">LinkedIn</a></p><p>Electrical Engineering &amp; Technical Support<br>© 2026 Eric John Cruz</p></div></footer></body></html>'''
    if filename == '404.html' and BASE:
        # Pages serves this document at arbitrary missing paths, including nested URLs.
        out = re.sub(r'((?:href|src)=")([^"#][^"]*)"',
                     lambda m: m.group(0) if re.match(r'[a-zA-Z][a-zA-Z0-9+.-]*:|//',m.group(2))
                     else m.group(1)+esc(BASE)+'/'+m.group(2)+'"',out)
    (OUT/filename).write_text(out,encoding='utf-8')

def pagehead(kicker,title,intro):return f'<div class="page-head"><span class="eyebrow">{esc(kicker)}</span><h1>{esc(title)}</h1><p class="intro">{esc(intro)}</p></div>'
def project_card(c): return f'''<article class="card project"><div class="visual">{c['visual']()}</div><div class="card-body"><span class="tag">{esc(c['classification'])} · {esc(c['category'])}</span><h3><a href="case-{c['slug']}.html">{esc(c['title'])}</a></h3><p>{esc(c['intro'])}</p><a class="text-link" href="case-{c['slug']}.html">View worksheet and assumptions</a></div></article>'''

PATHWAYS = [
 ('engineering','01','Electrical engineering and estimating','Electrical layouts, load and panel schedules, quantity takeoffs and bills of quantities.'),
 ('technical-sales','02','Technical sales and applications','Equipment specifications, supplier questions and quotation information, supported by customer communication experience.'),
 ('operations','03','Engineering operations and remote support','Document registers, revisions, spreadsheet reports and follow-up on outstanding records.')
]

def portrait(priority=False):
    return '<img class="professional-portrait" src="assets/eric-john-cruz.webp" width="232" height="232" alt="Eric John Cruz, Electrical Engineer." decoding="async"'+(' fetchpriority="high"' if priority else ' loading="lazy"')+'>'

def render_site():
    home=f'''<section class="hero"><div><span class="eyebrow">Eric John Cruz · Philippines</span><h1>Electrical drawings, quantity takeoffs and technical records.</h1><p class="intro">I prepare electrical layouts, schedules and quantity takeoffs for small building projects. My experience also includes power plant electrical and automation records, supervised municipal engineering work, and spreadsheet reporting.</p><div class="actions">{btn('View work samples','projects.html')}{btn('View résumé','resumes.html',True)}</div><p class="location">Based in Puerto Princesa City · Open to engineering and technical support opportunities</p></div><aside class="hero-figure"><div class="label">From my work history</div><div class="hero-summary"><h2>Drawings, schedules and equipment records</h2><ul><li>Electrical layouts, single-line diagrams and material lists for small projects</li><li>Cable schedules, equipment registers and manuals from power plant documentation work</li><li>Inspection records and drafting during supervised municipal training</li></ul><a class="text-link" href="experience.html">See roles and dates</a></div></aside></section><div class="credentials"><div><small>Professional qualification</small>Registered Electrical Engineer (Philippines)</div><div><small>Engineering education</small>BS Electrical Engineering · Palawan State University</div><div><small>Professional background</small>Engineering documentation + remote operations</div></div>'''
    home=home.replace('<h1>', '<div class="hero-identity">'+portrait(True)+'<div><strong>Electrical Engineer</strong><p>Electrical Design | Estimating | Technical Documentation</p></div></div><h1>',1)
    home+='<section class="section"><div class="section-head"><h2>Areas of work</h2><p>Electrical drawings, quantity takeoffs, equipment information and technical records.</p></div><div class="grid">'
    for slug,num,title,desc in PATHWAYS:home+=f'<article class="card"><span class="number">{num}</span><h3>{title}</h3><p>{desc}</p><a class="text-link" href="{slug}.html">See details</a></article>'
    home+='''</div></section><section class="section"><div class="section-head"><h2>Engineering work I have done</h2><p>My work history covers drawings, engineering records and supervised site documentation. Project files and employer records are not displayed publicly.</p></div><div class="grid two"><article class="card"><h3>Independent electrical projects</h3><p>Prepared layouts, single-line diagrams, panel and load schedules, material lists and bills of quantities for small building projects.</p></article><article class="card"><h3>Power plant technical records</h3><p>Maintained equipment lists, cable and panel schedules, manuals and test records. Checked references and followed up on missing information with site engineers.</p></article><article class="card"><h3>Municipal engineering placement</h3><p>Assisted with street-lighting and building inspection records, AutoCAD drawings and cost estimates under supervision.</p></article></div><div class="actions"><a class="button secondary" href="experience.html">See the experience details</a></div></section>'''
    home+='<section class="section"><div class="section-head"><div><span class="eyebrow">Engineering work samples</span><h2>Sample calculations and technical records</h2></div><p>These samples demonstrate takeoff, calculation, equipment-review and document-control methods with their input assumptions shown.</p></div><div class="grid two">'+''.join(project_card(c) for c in CASES)+'</div></section>'
    home+='<section class="section"><div class="section-head"><div><span class="eyebrow">Ongoing professional training</span><h2>Australian Electrical Estimating</h2></div><p>Drawing and specification review, GroundPlan takeoffs, quantity schedules, allowances and supplier RFQ scope.</p></div><a class="text-link" href="training.html">Read about my current training</a></section>'
    home+='<section class="section"><div class="section-head"><h2>Résumé and contact</h2><p>View my engineering résumé or contact me about a relevant opening.</p></div><div class="actions">'+btn('View résumé','resumes.html')+btn('Contact Eric','contact.html',True)+'</div></section>'
    page('index.html','Electrical Engineering and Documentation','Eric John Cruz, Philippine Registered Electrical Engineer. Building electrical drawings, schedules, engineering records and quantity takeoffs.',home)

    exp=pagehead('Experience','Electrical drawings, plant records and reporting','I have worked on small-building electrical drawings and schedules, handled electrical and automation records in a power plant assignment, and assisted with municipal inspection documentation.')
    exp='<div class="experience-intro">'+exp+'<figure class="experience-portrait">'+portrait()+'<figcaption>Eric John Cruz, REE</figcaption></figure></div>'
    exp+='<section><h2>Engineering experience</h2><div class="timeline">'
    def job(e):return f'<article class="job"><div class="date">{esc(e["dates"])}<br>{esc(e["type"])}</div><div><h3>{esc(e["title"])}</h3><p class="company">{esc(e["company"])}</p>{ul(e["bullets"])}</div></article>'
    exp+=''.join(job(e) for e in PROFILE['experience'][:3])+'</div></section><section class="section"><h2>Remote business and customer support</h2><p class="intro">This is separate from my electrical engineering work, but involved reporting, record keeping and following up on outstanding tasks.</p>'+''.join(job(e) for e in PROFILE['experience'][3:])+'</section>'
    exp+='<section class="section"><h2>Education and professional development</h2><p><strong>'+PROFILE['education']+'</strong><br>'+PROFILE['university']+' · '+PROFILE['education_dates']+'</p><p><strong>Registered Electrical Engineer (Philippines)</strong></p><p>Ongoing Australian electrical estimating training covers drawing review, GroundPlan takeoffs, quantity schedules, assumptions, supplier packages and switchboard RFQ scope.</p><a class="text-link" href="training.html">See my ongoing estimating training</a></section>'
    exp+='<section class="section"><h2>Tools in context</h2><dl class="tool-list">'+''.join(f'<div><dt>{esc(t["name"])}</dt><dd>{esc(t["context"])}</dd></div>' for t in PROFILE['tools'])+'</dl></section>'
    page('experience.html','Professional Experience','Engineering, documentation, customer support and remote professional experience of Eric John Cruz.',exp,'experience')

    configs={
     'engineering':('Electrical engineering and estimating','Electrical drawings and quantity takeoffs','I prepare electrical drawings, schedules and bills of quantities for small building projects. My current training covers Australian electrical estimating and GroundPlan quantity takeoffs.',
      [('Drawing and scope review','Read electrical layouts, legends, schedules and scope notes. Identify missing information and carry it into a clarification register.'),('Takeoffs and BOQs','Organize device counts and material quantities with drawing references, units, assumptions and revision notes.'),('Design documentation support','Prepare or revise electrical layouts, single-line diagrams and load/panel schedules for review within the agreed project scope.'),('Estimating development','Continue GroundPlan practice and structured tender review. GroundPlan activities are part of ongoing professional training.')],['takeoff','schedule-review'],'engineering'),
     'technical-sales':('Technical sales and applications','Equipment research and quotation support','I am pursuing technical sales and applications support, drawing on my engineering studies, documentation work and customer service experience.',
      [('Requirements clarification','Identify the specifications, quantities and unanswered questions in a customer request.'),('Equipment research and comparison','Compare supplied specifications and identify missing data, assumptions, alternatives and manufacturer confirmations.'),('Quotation support','Prepare scope summaries, schedules and supplier questions. Track inclusions, exclusions, validity and follow-up actions.'),('Customer communication','Use my customer support and technical documentation background when explaining product information or following up on quotations.')],['equipment-evaluation','takeoff'],'technical-sales'),
     'operations':('Engineering operations and remote support','Document registers and project follow-up','My power plant document-control and remote reporting experience applies to technical administration and project support.',
      [('Document control','Track document numbers, revisions, missing attachments and review status.'),('Project coordination','Track outstanding documents, actions and due dates, and follow up when information is missing.'),('Research and reporting','Prepare spreadsheets and status reports from documents and project records.'),('Technical administration','Organize technical schedules, review comments, outstanding information and handover records.')],['document-control','equipment-evaluation'],'operations')
    }
    for slug,(kicker,title,intro,blocks,selected,variant) in configs.items():
        body=pagehead(kicker,title,intro)+'<section class="section"><div class="grid two">'+''.join(f'<article class="card"><h2 style="font-size:1.5rem">{esc(h)}</h2><p>{esc(t)}</p></article>' for h,t in blocks)+'</div></section>'
        if slug=='technical-sales':body+='<p class="notice">I am applying for technical sales and applications support roles. My work history is in engineering documentation and customer support, not a previous sales engineer position.</p>'
        if slug=='engineering':body+='<p class="notice">Engineering work follows the agreed scope and applicable review requirements. My REE credential is a Philippine qualification; overseas regulated design and approval require the relevant local authority.</p>'
        body+='<section class="section"><h2>Related work samples</h2><div class="grid two">'+''.join(project_card(c) for c in CASES if c['slug'] in selected)+'</div></section><div class="actions">'+btn('View résumé','resume.html')+btn('Contact me','contact.html',True)+'</div>'
        page(slug+'.html',kicker,intro,body,'expertise')

    body=pagehead('Projects and work samples','Engineering work samples','Takeoff schedules, electrical calculations, equipment evaluation and document registers with traceable inputs and assumptions.')
    body+='<h2>Sample projects and calculations</h2><div class="grid two">'+''.join(project_card(c) for c in CASES)+'</div><section class="section"><article class="card"><span class="tag">Ongoing professional training</span><h2>Australian Electrical Estimating</h2><p>Current training in drawing review, GroundPlan takeoffs, quantity schedules, assumptions, supplier packages and switchboard RFQ scope.</p><a class="text-link" href="training.html">Read the ongoing training account</a></article></section>'
    body+='<section class="section prose" id="sample-basis"><h2>Sample basis and intended use</h2>'+p(SAMPLE_PROVENANCE)+p('Each worksheet identifies its assumed quantities, ratings and review conditions. The samples demonstrate methods without reproducing employer records or proprietary training drawings.')+'<a class="text-link" href="experience.html">Read my professional work history</a></section>'
    page('projects.html','Engineering Work Samples','Sample electrical estimating, connected-load calculation, equipment evaluation and engineering document registers, with assumptions shown.',body,'projects')

    for c in CASES:
        body='<div class="page-head"><p class="breadcrumb"><a href="projects.html">Work samples</a> / '+esc(c['category'])+'</p><span class="eyebrow">'+esc(c['classification'])+'</span><h1>'+esc(c['title'])+'</h1><p class="intro">'+esc(c['intro'])+'</p></div>'
        body+='<div class="content-grid"><article class="prose"><figure class="figure">'+c['visual']()+'<figcaption>'+esc(c['caption'])+'</figcaption></figure>'
        for key,label in [('context','Sample inputs'),('objective','Objective'),('approach','Technical approach')]:body+='<h2>'+label+'</h2>'+p(c[key])
        body+=details(c['slug'])
        for key,label in [('deliverables','Sample outputs'),('checks','Calculation and consistency checks'),('result','Illustrative result'),('limit','Assumptions and limits')]:body+='<h2>'+label+'</h2>'+p(c[key])
        body+='<p class="small"><a href="projects.html#sample-basis">Sample origin and intended use</a></p></article><aside class="aside"><h2>At a glance</h2><dl><dt>Focus</dt><dd>'+esc(c['category'])+'</dd><dt>Tools</dt><dd>'+esc(c['tools'])+'</dd></dl><div class="actions">'+btn('Open the sample sheet','sample-'+c['slug']+'.html')+f'<a class="button secondary" href="samples/{c["slug"]}.json" download>Download the example data</a></div></aside></div>'
        page('case-'+c['slug']+'.html',c['title'],c['intro'],body,'projects')
        sample=pagehead('Sample worksheet',c['title'],c['context'])+'<h2>Inputs and outputs</h2>'+details(c['slug'])
        if c['slug']=='takeoff':sample+=table(['Point ID','Type','Drawing X','Drawing Y'],[[v['id'],v['type'],v['x'],v['y']] for v in DEMO['points']],'Point register in SVG drawing coordinates')
        sample+='<div class="prose"><h2>Checks</h2>'+p(c['checks'])+'<h2>Assumptions and limits</h2>'+p(c['limit'])+'<p class="small"><a href="projects.html#sample-basis">Sample origin and intended use</a></p></div><div class="actions">'+btn('Return to the case study','case-'+c['slug']+'.html')+'</div>'
        page('sample-'+c['slug']+'.html',c['title']+' Sample',c['intro'],sample,'projects')

    train=pagehead('Ongoing professional training','Australian Electrical Estimating','I am currently learning and practicing Australian electrical quantity-takeoff methods, drawing review and supplier-scope checks.')
    train+='<div class="content-grid"><article class="prose">'+''.join('<h2>'+h+'</h2>'+p(t) for h,t in [
      ('Context','An Australian residential estimating training project involves multiple electrical systems, general notes and supplier-quoted packages. The public account describes the method without publishing the project drawings, specification or course material.'),
      ('Objective','Practice reading electrical drawings, legends, specifications and general notes; organize takeoff schedules with measured quantities, allowances and scope questions clearly identified.'),
      ('My role','Review drawings and notes, organize system-based takeoff items, and raise clarification questions during training. These activities are part of ongoing professional development; they are separate from client estimates.'),
      ('Technical approach','Start with the drawing register and known scale references. Read notes and legends before counting. Label an allowance in the item name and carry the fuller explanation into the assumptions register. Keep supplier package inclusions visible when reviewing quotations.'),
      ('Tools','GroundPlan for practice measurements and counts; spreadsheets and written registers for assumptions, revisions and scope questions.'),
      ('Deliverables and checks','The activities include GroundPlan counts and lengths, quantity schedules, note-based allowances, communications questions, supplier-quoted package inclusions and switchboard RFQ scope. In training, I keep the drawing or note reference with each item and record questions for instructor review.'),
      ('Current status','Training is ongoing. It does not represent a completed Australian estimating contract, formal certification, Australian electrical licensure or an approved final estimate.')])+'</article><aside class="aside"><h2>Training status</h2><p class="small">Ongoing professional training. Original drawings, specifications and course materials remain private.</p><div class="actions">'+btn('See the sample takeoff method','case-takeoff.html')+'</div></aside></div>'
    page('training.html','Australian Electrical Estimating Training','Ongoing training in Australian drawing review, GroundPlan takeoffs, assumptions, supplier packages and switchboard RFQ scope.',train,'projects')

    contact=pagehead('Contact','Contact Eric John Cruz','Email me about an electrical engineering, estimating or technical support position.')
    contact+='<div class="grid two"><section class="contact-block"><h2>Email Eric</h2><p><a class="wide-link" href="mailto:'+PROFILE['email']+'?subject=Engineering%20or%20technical%20support%20opportunity">'+PROFILE['email']+'</a></p><button class="copy" data-copy-email type="button">Copy email address</button><p class="small" id="copy-status" aria-live="polite"></p><p class="small">Puerto Princesa City, Palawan, Philippines</p></section><section class="contact-block"><h2>Useful details to include</h2>'+ul(['Role or project focus','Type of drawings, records or technical information involved','Expected deliverables and review process','Location, working arrangement and target timing'])+'<a class="text-link" href="'+PROFILE['linkedin']+'">View my LinkedIn profile</a></section></div><div class="actions">'+btn('View résumé','resumes.html')+btn('View work samples','projects.html',True)+'</div>'
    page('contact.html','Contact','Contact Eric John Cruz about engineering, estimating, technical support and remote operations opportunities.',contact,'contact')

VARIANTS={
 'engineering':{
  'label':PROFILE['professional_title'],
  'filename':'Eric_John_Cruz_Engineering_Resume',
  'summary':'Registered Electrical Engineer (Philippines) with experience preparing electrical layouts, single-line diagrams, load and panel schedules, takeoffs and bills of quantities for small building projects. Background includes power plant electrical and automation documentation, supervised municipal engineering work, and technical records review.',
  'skills':[
    'Electrical drawings and schedules: layouts, single-line diagrams, load and panel schedules.',
    'Estimating documentation: material takeoffs, bills of quantities, scope notes and drawing revisions.',
    'Technical records: equipment lists, cable schedules, manuals, test documents and revision registers.',
    'Tools used for drafting, review and records: AutoCAD 2D, Bluebeam Revu, Microsoft Excel, Google Sheets, Microsoft Office and Google Workspace. GroundPlan: training and practice.'
  ]}
}

def resume_sections(key='engineering'):
    v=VARIANTS[key]
    items=[('h2','Professional summary'),('p',v['summary']),('h2','Technical skills')]
    items += [('bullet',x) for x in v['skills']]
    items += [('h2','Engineering experience')]
    for id in ['freelance','rapid','kalayaan']:
        e=next(x for x in PROFILE['experience'] if x['id']==id)
        items += [('h3',e['title']),('meta',e['company']+' | '+e['dates'])]
        items += [('bullet',x) for x in e['bullets']]
    items += [('break',''),('h2','Additional professional experience')]
    e=next(x for x in PROFILE['experience'] if x['id']=='remote')
    items += [('h3',e['title']),('meta',e['company']+' | '+e['dates'])]
    items += [('bullet','Prepared spreadsheet reports, research notes and client updates for remote digital marketing engagements, separate from electrical engineering assignments.'),('bullet','Tracked tasks and review comments, checked deliverables and followed up on missing information with remote team members.')]
    e=next(x for x in PROFILE['experience'] if x['id']=='foundever')
    items += [('h3',e['title']),('meta',e['company']+' | '+e['dates'])]+[('bullet',x) for x in e['bullets']]
    items += [('h2','Education and professional qualification'),('p',PROFILE['education']+' | '+PROFILE['university']+' | '+PROFILE['education_dates']),('p',PROFILE['credential'])]
    items += [('h2','Australian electrical estimating training'),('p','Ongoing professional development in drawing, legend and specification review; GroundPlan takeoffs; quantity schedules; assumptions and allowances; supplier package and switchboard RFQ scope review. These activities are training, separate from completed client estimates.')]
    if BASE:items += [('h2','Portfolio'),('p',BASE+'/')]
    return items

def resume_html(key,sections):
    v=VARIANTS[key]
    s='<article class="resume-paper"><h1>'+PROFILE['name']+'</h1><p><strong>'+v['label']+'</strong></p><p class="resume-meta">'+PROFILE['location']+'<br><a href="mailto:'+PROFILE['email']+'">'+PROFILE['email']+'</a><br><a href="'+PROFILE['linkedin']+'">linkedin.com/in/eric-john-cruz-ree/</a></p>'
    inlist=False
    for kind,text in sections:
        if kind=='bullet':
            if not inlist:s+='<ul>';inlist=True
            s+='<li>'+esc(text)+'</li>';continue
        if inlist:s+='</ul>';inlist=False
        if kind=='break':s+='<div class="page-break"></div>'
        elif kind in ['h2','h3']:s+=f'<{kind}>{esc(text)}</{kind}>'
        else:
            content='<a href="'+esc(text)+'">'+esc(text)+'</a>' if text.startswith('https://') else esc(text)
            s+=f'<p{" class="+chr(34)+"resume-meta"+chr(34) if kind=="meta" else ""}>{content}</p>'
    if inlist:s+='</ul>'
    return s+'</article>'

def word_link(paragraph,text,url,size=9):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    link=OxmlElement('w:hyperlink')
    link.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');properties=OxmlElement('w:rPr')
    font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Arial');font.set(qn('w:hAnsi'),'Arial');properties.append(font)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));properties.append(sz)
    color=OxmlElement('w:color');color.set(qn('w:val'),'000000');properties.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');properties.append(underline)
    run.append(properties);t=OxmlElement('w:t');t.text=text;run.append(t);link.append(run);paragraph._p.append(link)

def create_docx(key,sections):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.58);sec.left_margin=sec.right_margin=Inches(.66)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
        st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    # Remove inherited title rules from the default Word template.
    for st in d.styles:
        ppr=st.element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
    normal=d.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.08
    d.styles['Title'].font.size=Pt(23);d.styles['Title'].paragraph_format.space_after=Pt(6)
    for name,size in [('Heading 1',11),('Heading 2',10.5)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(11);st.paragraph_format.space_after=Pt(5)
    d.styles['List Bullet'].paragraph_format.space_after=Pt(4)
    d.core_properties.author=PROFILE['name'];d.core_properties.title=VARIANTS[key]['label'];d.core_properties.subject='Professional resume';d.core_properties.keywords=''
    d.add_paragraph(PROFILE['name'],style='Title')
    q=d.add_paragraph();q.add_run(VARIANTS[key]['label']).bold=True
    for x,url in [(PROFILE['location'],None),(PROFILE['email'],'mailto:'+PROFILE['email']),(PROFILE['linkedin'],PROFILE['linkedin'])]:
        q=d.add_paragraph();q.paragraph_format.space_after=Pt(3)
        if url:word_link(q,x,url)
        else:q.add_run(x).font.size=Pt(9)
    for kind,text in sections:
        if kind=='break':d.add_page_break();continue
        if kind=='h2':q=d.add_paragraph(text.upper(),style='Heading 1')
        elif kind=='h3':q=d.add_paragraph(text,style='Heading 2')
        elif kind=='bullet':q=d.add_paragraph(text,style='List Bullet')
        else:
            q=d.add_paragraph()
            if kind=='p' and text.startswith('https://'):word_link(q,text,text,10.5)
            else:q.add_run(text)
            if kind=='meta':
                for r in q.runs:r.font.size=Pt(9)
                q.paragraph_format.keep_with_next=True
        q.paragraph_format.widow_control=True
    footer=sec.footer.paragraphs[0];footer.alignment=2
    footer.add_run(PROFILE['name']+' | ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for r in footer.runs:r.font.size=Pt(8)
    d.save(OUT/'downloads'/(VARIANTS[key]['filename']+'.docx'))

def create_pdf(key,sections):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Name',fontName='Helvetica-Bold',fontSize=23,leading=26,spaceAfter=9))
    styles.add(ParagraphStyle(name='Role',fontName='Helvetica-Bold',fontSize=11,leading=14,spaceAfter=7))
    styles.add(ParagraphStyle(name='Contact',fontName='Helvetica',fontSize=9,leading=12,spaceAfter=2))
    styles.add(ParagraphStyle(name='Section',fontName='Helvetica-Bold',fontSize=11,leading=14,spaceBefore=12,spaceAfter=5,keepWithNext=True))
    styles.add(ParagraphStyle(name='Job',fontName='Helvetica-Bold',fontSize=10.5,leading=13,spaceBefore=8,spaceAfter=3,keepWithNext=True))
    styles.add(ParagraphStyle(name='Meta',fontName='Helvetica',fontSize=9,leading=12,spaceAfter=4,keepWithNext=True))
    styles.add(ParagraphStyle(name='Copy',fontName='Helvetica',fontSize=10.5,leading=13.2,spaceAfter=5))
    styles.add(ParagraphStyle(name='BulletText',parent=styles['Copy'],leftIndent=11,firstLineIndent=-9,spaceAfter=4))
    doc=SimpleDocTemplate(str(OUT/'downloads'/(VARIANTS[key]['filename']+'.pdf')),pagesize=letter,leftMargin=47.5,rightMargin=47.5,topMargin=41.8,bottomMargin=41.8,title=VARIANTS[key]['label'],author=PROFILE['name'])
    story=[Paragraph(esc(PROFILE['name']),styles['Name']),Paragraph(esc(VARIANTS[key]['label']),styles['Role'])]
    for x,url in [(PROFILE['location'],None),(PROFILE['email'],'mailto:'+PROFILE['email']),(PROFILE['linkedin'],PROFILE['linkedin'])]:story.append(Paragraph('<link href="'+esc(url)+'">'+esc(x)+'</link>' if url else esc(x),styles['Contact']))
    for kind,text in sections:
        if kind=='break':story.append(PageBreak());continue
        style={'h2':'Section','h3':'Job','meta':'Meta','p':'Copy','bullet':'BulletText'}[kind]
        text=text.upper() if kind=='h2' else text
        content='<link href="'+esc(text)+'">'+esc(text)+'</link>' if kind=='p' and text.startswith('https://') else esc(text)
        story.append(Paragraph(('• ' if kind=='bullet' else '')+content,styles[style]))
    def footer(canvas,doc):
        canvas.setFont('Helvetica',8);canvas.drawRightString(564,24,f'{PROFILE["name"]} | {doc.page}')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)

def render_resumes():
    key='engineering';v=VARIANTS[key];sections=resume_sections(key)
    create_docx(key,sections);create_pdf(key,sections)
    links=btn('Download PDF','downloads/'+v['filename']+'.pdf')+btn('Download Word','downloads/'+v['filename']+'.docx',True)
    hub=pagehead('Résumé','Engineering résumé','Electrical drawings, estimating documentation, technical records and engineering support experience.')
    hub+='<section class="resume-hub"><h2>Eric John Cruz</h2><p>'+esc(v['label'])+'</p><p class="small">Registered Electrical Engineer (Philippines)</p><div class="actions">'+btn('View Résumé','resume.html')+'<a class="button secondary" href="downloads/'+v['filename']+'.pdf" download>Download PDF</a><a class="button secondary" href="downloads/'+v['filename']+'.docx" download>Download Word</a></div></section>'
    page('resumes.html','Engineering Résumé','View Eric John Cruz’s general engineering résumé and download the PDF or Word document.',hub,'resumes')
    body='<div class="actions no-print">'+links+'</div>'+resume_html(key,sections)
    page('resume.html','Engineering Résumé',v['summary'],body,'resumes')

def main():
    global BASE
    parser=argparse.ArgumentParser();parser.add_argument('--base-url',default='');args=parser.parse_args();BASE=args.base_url.rstrip('/')
    if BASE and not re.fullmatch(r'https://[A-Za-z0-9.-]+(?:/[A-Za-z0-9._~-]+)*',BASE):raise ValueError('Use an HTTPS base URL with no query or fragment.')
    if OUT.exists():
        if OUT.resolve().parent!=ROOT.resolve() or OUT.name!='dist':raise ValueError('Output directory must remain inside the repository.')
        shutil.rmtree(OUT)
    for folder in ['assets','downloads','samples']:(OUT/folder).mkdir(parents=True,exist_ok=True)
    for file in (ROOT/'assets').iterdir():shutil.copyfile(file,OUT/'assets'/file.name)
    (OUT/'assets/favicon.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64"><rect width="64" height="64" rx="32" fill="#122c35"/><text x="32" y="41" text-anchor="middle" font-family="Georgia,serif" font-size="29" fill="white">EC</text></svg>',encoding='utf-8')
    data={'takeoff':{'points':DEMO['points'],'route':DEMO['route']},'schedule-review':{'voltage_v':230,'loads':DEMO['loads']},'equipment-evaluation':DEMO['supply'],'document-control':DEMO['records']}
    for slug,value in data.items():
        case=next(c for c in CASES if c['slug']==slug)
        sample={'classification':case['title'],'provenance':SAMPLE_PROVENANCE,'input_basis':SAMPLE_BASIS[slug],'inputs':value,'checks':case['checks'],'limitations':case['limit']}
        (OUT/'samples'/(slug+'.json')).write_text(json.dumps(sample,indent=2),encoding='utf-8')
    render_site();render_resumes()
    page('404.html','Page Not Found','This page could not be found.',pagehead('Page not found','Page not found','Use the navigation to find work samples, experience or the résumé.')+btn('Open the homepage','index.html'))
    (OUT/'.nojekyll').write_text('')
    if BASE:
        urls=[BASE+'/'+x.name for x in sorted(OUT.glob('*.html')) if x.name not in ['404.html'] and not x.name.startswith('sample-')]
        (OUT/'sitemap.xml').write_text('<?xml version="1.0" encoding="UTF-8"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'+''.join('<url><loc>'+esc(u)+'</loc></url>' for u in urls)+'</urlset>',encoding='utf-8')
        (OUT/'robots.txt').write_text('User-agent: *\nAllow: /\nSitemap: '+BASE+'/sitemap.xml\n')
    else:
        (OUT/'robots.txt').write_text('User-agent: *\nDisallow: /\n')
        if (OUT/'sitemap.xml').exists():(OUT/'sitemap.xml').unlink()
    (ROOT/'qa').mkdir(exist_ok=True)
    (ROOT/'qa/build.json').write_text(json.dumps({'base_url':BASE,'pages':len(list(OUT.glob('*.html'))),'resume_variants':list(VARIANTS),'status':'Built locally; deployment verification separate'},indent=2))
    print(f'Built {len(list(OUT.glob("*.html")))} pages, 1 DOCX résumé, 1 PDF and 4 sample datasets.')

if __name__=='__main__':main()
