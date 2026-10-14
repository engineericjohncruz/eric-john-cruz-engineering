"""Content, arithmetic, local link and document validation."""
from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urlsplit,unquote
from zipfile import ZipFile
from pypdf import PdfReader
from docx import Document
import json, math, sys
import build
from build import OUT,ROOT,DEMO,VARIANTS,resume_sections,record_ready

build.BASE=json.loads((ROOT/'qa/build.json').read_text(encoding='utf-8'))['base_url']
assert len(VARIANTS)==1,'Only one public resume configuration is allowed'
assert build.PROFILE['credential']=='Registered Electrical Engineer (Philippines)'

class Page(HTMLParser):
    def __init__(self):super().__init__();self.links=[];self.ids=set();self.h1=0;self.title=False;self.description=False;self.errors=[];self.lang=False
    def handle_starttag(self,tag,attrs):
        a=dict(attrs)
        if tag=='html':self.lang=a.get('lang')=='en'
        if tag=='h1':self.h1+=1
        if tag=='title':self.title=True
        if tag=='meta' and a.get('name')=='description':self.description=bool(a.get('content'))
        if 'id' in a:
            if a['id'] in self.ids:self.errors.append('Duplicate ID: '+a['id'])
            self.ids.add(a['id'])
        if tag=='img' and 'alt' not in a:self.errors.append('Image missing alt')
        for key in ['href','src']:
            if key in a:self.links.append(a[key])

errors=[];pages={};checked_links=0
for f in OUT.glob('*.html'):
    parser=Page();text=f.read_text(encoding='utf-8');parser.feed(text);pages[f]=parser
    if parser.h1!=1:errors.append(f'{f.name}: expected one h1, found {parser.h1}')
    if not(parser.title and parser.description and parser.lang):errors.append(f'{f.name}: metadata missing')
    errors += [f'{f.name}: {e}' for e in parser.errors]
    for forbidden in ['+63 938','libfile_','file_0000','1dzqw_U','salary','commissioned for','100% compliant']:
        if forbidden in text:errors.append(f'{f.name}: forbidden/private content {forbidden}')
for f,parser in pages.items():
    for link in parser.links:
        if build.BASE and link.startswith(build.BASE+'/'):
            link=link[len(build.BASE)+1:] or 'index.html'
        u=urlsplit(link)
        if u.scheme or link.startswith('//'):continue
        target=(f.parent/unquote(u.path)).resolve() if u.path else f.resolve()
        if not target.is_relative_to(OUT.resolve()):errors.append(f'{f.name}: path escapes output: {link}')
        if not target.exists():errors.append(f'{f.name}: missing {link}')
        if u.fragment and target in pages and u.fragment not in pages[target].ids:errors.append(f'{f.name}: unknown anchor {link}')
        checked_links+=1

if build.BASE:
    not_found=(OUT/'404.html').read_text(encoding='utf-8')
    for path in ['assets/site.css','assets/site.js','index.html']:
        if '"'+build.BASE+'/'+path+'"' not in not_found:
            errors.append('404.html: missing absolute recovery URL for '+path)

assert len(DEMO['points'])==25
assert len({x['id'] for x in DEMO['points']})==25
assert sum(x['type']=='lighting' for x in DEMO['points'])==12
assert sum(x['type']=='outlets' for x in DEMO['points'])==13
assert sum(DEMO['route']['segments_m'])==39
assert sum(DEMO['route']['segments_m']) + DEMO['route']['rise_and_termination_allowance_m'] == 45
assert 'procurement_allowance_fraction' not in DEMO['route']
assert sum(x['quantity']*x['unit_va'] for x in DEMO['loads'])==1794
assert math.isclose(1794/230,7.8)
target=sum(DEMO['supply']['load_a'])*(1+DEMO['supply']['reserve_fraction'])
assert math.isclose(target,1.875)
assert [x['name'] for x in DEMO['supply']['options'] if x['available_a']>=target]==['B']
assert sum(not record_ready(r) for r in DEMO['records'])==3
assert not record_ready(dict(DEMO['records'][0],attachment=False))
assert not record_ready(dict(DEMO['records'][0],review='Pending'))
assert not record_ready(dict(DEMO['records'][0],received='B'))

expected_downloads={'Eric_John_Cruz_Engineering_Resume.pdf','Eric_John_Cruz_Engineering_Resume.docx'}
assert {f.name for f in (OUT/'downloads').iterdir()}==expected_downloads,'Unexpected public resume files'
for old in ['resume-engineering.html','resume-technical-sales.html','resume-operations.html']:
    assert not (OUT/old).exists(),'Obsolete resume page remains'
    for html in OUT.glob('*.html'):
        assert old not in html.read_text(encoding='utf-8'),'Obsolete resume link remains'
    if (OUT/'sitemap.xml').exists():assert old not in (OUT/'sitemap.xml').read_text()
for html in OUT.glob('*.html'):
    body=html.read_text(encoding='utf-8')
    assert '<small>Electrical Engineering &amp; Technical Support</small>' in body
    assert 'Choose a résumé' not in body and 'All résumé versions' not in body and '>Résumés<' not in body
docs=[]
def normalize(s):return ' '.join(s.replace('•',' ').split())
for key,v in VARIANTS.items():
    pdf=OUT/'downloads'/(v['filename']+'.pdf');word=OUT/'downloads'/(v['filename']+'.docx')
    r=PdfReader(pdf);text=' '.join(p.extract_text() or '' for p in r.pages)
    wd=Document(word);wt=' '.join(p.text for p in wd.paragraphs)
    if len(r.pages)!=2:errors.append(f'{pdf.name}: {len(r.pages)} pages, expected 2')
    for kind,value in resume_sections(key):
        if kind=='break':continue
        if normalize(value).lower() not in normalize(text).lower():errors.append(f'{pdf.name}: missing content {value[:45]}')
        if normalize(value).lower() not in normalize(wt).lower():errors.append(f'{word.name}: missing content {value[:45]}')
    with ZipFile(word) as z:
        assert z.testzip() is None
        assert 'word/document.xml' in z.namelist()
        assert not any(n.startswith('word/media/') for n in z.namelist()),'Resume must be photo-free'
        relationships=z.read('word/_rels/document.xml.rels').decode()
        for url in ['mailto:'+build.PROFILE['email'],build.PROFILE['linkedin']]+([build.BASE+'/'] if build.BASE else []):
            assert url in relationships,'Missing Word contact link'
    pdf_links=[str(a.get_object().get('/A',{}).get('/URI','')) for page in r.pages for a in page.get('/Annots',[])]
    for url in ['mailto:'+build.PROFILE['email'],build.PROFILE['linkedin']]+([build.BASE+'/'] if build.BASE else []):
        assert url in pdf_links,'Missing PDF contact link'
    assert not wd.tables,'Resume must have no layout tables'
    docs.append({'variant':key,'pdf_pages':len(r.pages),'pdf_text_characters':len(text),'docx_paragraphs':len(wd.paragraphs)})
report={'pages':len(pages),'local_links_checked':checked_links,'arithmetic':'passed','document_rules':'passed' if not errors else 'failed','documents':docs,'errors':errors}
(ROOT/'qa/content-checks.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report,indent=2));sys.exit(1 if errors else 0)
